from __future__ import annotations

import io
import json
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import PurePath
from typing import Iterable

from docx import Document
from PIL import Image
from pypdf import PdfReader
import pytesseract

from .models import Segment


class ExtractionError(ValueError):
    pass


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[tuple[str, str]] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(("visible_text", data))

    def handle_comment(self, data: str) -> None:
        if data.strip():
            self.parts.append(("html_comment", data))

    def handle_starttag(
        self, _tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        for name, value in attrs:
            if value and name.lower() in {"alt", "title", "aria-label", "href", "src", "style"}:
                self.parts.append((f"attribute:{name.lower()}", value))


def _segments_from_text(text: str, source: str, prefix: str) -> list[Segment]:
    paragraphs = [part.strip() for part in text.splitlines() if part.strip()]
    if not paragraphs and text.strip():
        paragraphs = [text.strip()]
    return [
        Segment(id=f"{prefix}-{index}", text=part, source=source)
        for index, part in enumerate(paragraphs, start=1)
    ]


def _extract_html(data: bytes, source: str) -> list[Segment]:
    parser = _HTMLTextExtractor()
    parser.feed(data.decode("utf-8", errors="replace"))
    return [
        Segment(
            id=f"html-{index}",
            text=text.strip(),
            source=source,
            location={"kind": kind},
        )
        for index, (kind, text) in enumerate(parser.parts, start=1)
        if text.strip()
    ]


def _extract_email(data: bytes, source: str) -> list[Segment]:
    message = BytesParser(policy=policy.default).parsebytes(data)
    segments: list[Segment] = []
    for header in ("subject", "from", "to", "reply-to"):
        value = message.get(header)
        if value:
            segments.append(
                Segment(
                    id=f"email-header-{header}",
                    text=str(value),
                    source=source,
                    location={"header": header},
                )
            )
    part_index = 0
    for part in message.walk():
        if part.is_multipart():
            continue
        part_index += 1
        payload = part.get_payload(decode=True) or b""
        content_type = part.get_content_type()
        if content_type == "text/html":
            for segment in _extract_html(payload, source):
                segment.id = f"email-{part_index}-{segment.id}"
                segments.append(segment)
        elif content_type.startswith("text/"):
            charset = part.get_content_charset() or "utf-8"
            segments.extend(
                _segments_from_text(
                    payload.decode(charset, errors="replace"),
                    source,
                    f"email-{part_index}",
                )
            )
    return segments


def _extract_pdf(data: bytes, source: str) -> list[Segment]:
    try:
        reader = PdfReader(io.BytesIO(data))
        segments: list[Segment] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                segments.append(
                    Segment(
                        id=f"pdf-page-{index}",
                        text=text,
                        source=source,
                        location={"page": index},
                    )
                )
        return segments
    except Exception as exc:
        raise ExtractionError(f"PDF extraction failed: {exc}") from exc


def _extract_docx(data: bytes, source: str) -> list[Segment]:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Word extraction failed: {exc}") from exc
    texts: list[tuple[str, str]] = []
    texts.extend(("paragraph", paragraph.text) for paragraph in document.paragraphs)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                texts.append(
                    (f"table-{table_index}-{row_index}-{cell_index}", cell.text)
                )
    for section_index, section in enumerate(document.sections, start=1):
        texts.extend(
            (f"header-{section_index}", paragraph.text)
            for paragraph in section.header.paragraphs
        )
        texts.extend(
            (f"footer-{section_index}", paragraph.text)
            for paragraph in section.footer.paragraphs
        )
    return [
        Segment(
            id=f"docx-{index}",
            text=text.strip(),
            source=source,
            location={"kind": kind},
        )
        for index, (kind, text) in enumerate(texts, start=1)
        if text.strip()
    ]


def _extract_image(data: bytes, source: str) -> list[Segment]:
    try:
        image = Image.open(io.BytesIO(data))
        result = pytesseract.image_to_data(
            image,
            output_type=pytesseract.Output.DICT,
        )
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError(
            "Image inspection requires the Tesseract OCR executable."
        ) from exc
    except Exception as exc:
        raise ExtractionError(f"Image OCR failed: {exc}") from exc

    lines: dict[tuple[int, int, int], list[str]] = {}
    confidences: dict[tuple[int, int, int], list[float]] = {}
    for index, text in enumerate(result["text"]):
        if not text.strip():
            continue
        key = (
            int(result["block_num"][index]),
            int(result["par_num"][index]),
            int(result["line_num"][index]),
        )
        lines.setdefault(key, []).append(text)
        try:
            confidence = float(result["conf"][index])
        except (TypeError, ValueError):
            confidence = -1
        confidences.setdefault(key, []).append(confidence)

    return [
        Segment(
            id=f"ocr-{index}",
            text=" ".join(words),
            source=source,
            location={
                "ocr_line": list(key),
                "mean_confidence": sum(confidences[key]) / len(confidences[key]),
            },
        )
        for index, (key, words) in enumerate(lines.items(), start=1)
    ]


def extract_segments(
    data: bytes,
    media_type: str,
    filename: str | None = None,
) -> list[Segment]:
    source = filename or media_type
    normalized_type = media_type.split(";", 1)[0].strip().lower()
    suffix = PurePath(filename or "").suffix.lower()

    if normalized_type == "text/html" or suffix in {".html", ".htm"}:
        segments = _extract_html(data, source)
    elif normalized_type == "message/rfc822" or suffix == ".eml":
        segments = _extract_email(data, source)
    elif normalized_type == "application/pdf" or suffix == ".pdf":
        segments = _extract_pdf(data, source)
    elif (
        normalized_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == ".docx"
    ):
        segments = _extract_docx(data, source)
    elif normalized_type.startswith("image/"):
        segments = _extract_image(data, source)
    elif normalized_type in {"application/json", "application/problem+json"}:
        try:
            parsed = json.loads(data)
        except json.JSONDecodeError as exc:
            raise ExtractionError(f"Invalid JSON: {exc}") from exc
        segments = _segments_from_text(
            json.dumps(parsed, ensure_ascii=False, indent=2),
            source,
            "json",
        )
    elif (
        normalized_type.startswith("text/")
        or suffix
        in {
            ".md",
            ".markdown",
            ".py",
            ".js",
            ".ts",
            ".tsx",
            ".jsx",
            ".java",
            ".cs",
            ".go",
            ".rs",
            ".sql",
            ".xml",
            ".yaml",
            ".yml",
        }
    ):
        segments = _segments_from_text(
            data.decode("utf-8", errors="replace"),
            source,
            "text",
        )
    else:
        raise ExtractionError(f"Unsupported media type: {media_type}")

    if not segments:
        raise ExtractionError("No inspectable text was extracted from the content.")
    return segments


def join_segments(segments: Iterable[Segment]) -> str:
    return "\n".join(segment.text for segment in segments)
